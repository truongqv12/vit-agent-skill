import os
import sys
import argparse
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from jira_api import JiraClient

def create_excel_report(issues, project_key, output_filename, base_url):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Tasks {project_key}"

    font_title = Font(name="Calibri", size=16, bold=True, color="003366")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="555555")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=11, color="000000")
    font_link = Font(name="Calibri", size=11, color="0000EE", underline="single")
    font_bold = Font(name="Calibri", size=11, bold=True)

    fill_header = PatternFill(start_color="004B8D", end_color="004B8D", fill_type="solid")
    fill_even = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    fill_status_done = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    fill_status_prog = PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid")
    fill_status_todo = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="D3D3D3")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_header = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style="medium", color="002B50"))

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 1. Title
    ws.merge_cells("A1:J1")
    ws["A1"] = f"DANH SÁCH TASK DỰ ÁN {project_key.upper()} (JIRA VNPT)"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:J2")
    ws["A2"] = f"Thời gian xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Tổng số: {len(issues)} tasks | Hệ thống: {base_url}"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 10

    # 2. Header
    headers = [
        ("STT", 6, align_center),
        ("Issue Key", 14, align_center),
        ("Loại", 12, align_center),
        ("Tóm tắt công việc (Summary)", 50, align_wrap_left),
        ("Trạng thái", 16, align_center),
        ("Người xử lý", 20, align_left),
        ("Người tạo", 20, align_left),
        ("Độ ưu tiên", 14, align_center),
        ("Hạn xử lý", 14, align_center),
        ("Ngày tạo", 14, align_center),
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 26

    for col_idx, (header_text, width, align) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # 3. Data rows
    start_row = 5
    for idx, issue in enumerate(issues, start=1):
        row_num = start_row + idx - 1
        ws.row_dimensions[row_num].height = 22
        fields = issue.get("fields", {})

        key = issue.get("key", "")
        summary = fields.get("summary", "")
        issuetype = fields.get("issuetype", {}).get("name", "") if fields.get("issuetype") else ""
        status_name = fields.get("status", {}).get("name", "") if fields.get("status") else ""
        priority = fields.get("priority", {}).get("name", "") if fields.get("priority") else ""
        assignee = fields.get("assignee", {}).get("displayName", "") if fields.get("assignee") else "Chưa gán"
        reporter = fields.get("reporter", {}).get("displayName", "") if fields.get("reporter") else ""
        duedate = fields.get("duedate", "") or ""
        created = fields.get("created", "")[:10] if fields.get("created") else ""

        row_fill = fill_even if idx % 2 == 0 else fill_odd

        values = [
            (idx, align_center, font_data, row_fill),
            (key, align_center, font_link, row_fill),
            (issuetype, align_center, font_data, row_fill),
            (summary, align_wrap_left, font_data, row_fill),
            (status_name, align_center, font_bold, row_fill),
            (assignee, align_left, font_data, row_fill),
            (reporter, align_left, font_data, row_fill),
            (priority, align_center, font_data, row_fill),
            (duedate, align_center, font_data, row_fill),
            (created, align_center, font_data, row_fill),
        ]

        for col_idx, (val, alignment, font, fill) in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = font
            cell.fill = fill
            cell.alignment = alignment
            cell.border = border_cell

            # Hyperlink for Jira key
            if col_idx == 2 and key:
                cell.hyperlink = f"{base_url}/browse/{key}"

            # Color highlight for status
            if col_idx == 5 and status_name:
                low_st = status_name.lower()
                if any(x in low_st for x in ["done", "resolved", "closed", "hoàn thành", "đã giải quyết"]):
                    cell.fill = fill_status_done
                elif any(x in low_st for x in ["in progress", "đang xử lý", "doing"]):
                    cell.fill = fill_status_prog
                else:
                    cell.fill = fill_status_todo

    ws.freeze_panes = "A5"
    wb.save(output_filename)
    print(f"[SUCCESS] Đã tạo thành công file Excel: {os.path.abspath(output_filename)}")

def main():
    parser = argparse.ArgumentParser(description="Export Jira project tasks to formatted Excel file")
    parser.add_argument("project", nargs="?", default="CCDT", help="Project key (default: CCDT)")
    parser.add_argument("-o", "--output", help="Output file name (.xlsx)")
    parser.add_argument("--unresolved", action="store_true", help="Only unresolved issues")
    parser.add_argument("--assignee", help="Filter by assignee username")
    args = parser.parse_args()

    client = JiraClient()
    proj = args.project.upper()
    jql_parts = [f'project = "{proj}"']
    if args.unresolved:
        jql_parts.append("resolution = Unresolved")
    if args.assignee:
        jql_parts.append(f'assignee = "{args.assignee}"')

    jql_query = " AND ".join(jql_parts) + " ORDER BY key ASC"
    output_file = args.output or f"Danh_sach_task_{proj}.xlsx"

    print("=" * 75)
    print(f"      XUẤT TOÀN BỘ TASK PROJECT {proj} RA FILE EXCEL")
    print("=" * 75)
    print(f" • JQL Query   : {jql_query}")
    print(f" • File đầu ra : {output_file}")
    print("\n[+] Đang tải dữ liệu từ Jira VNPT...")

    all_issues = []
    start_at = 0
    max_results = 100
    fields = "summary,status,issuetype,reporter,assignee,priority,created,duedate"

    while True:
        try:
            data = client.search_issues(jql_query, fields=fields, max_results=max_results, start_at=start_at)
            issues = data.get("issues", [])
            total = data.get("total", 0)
            all_issues.extend(issues)
            print(f" -> Đã tải {len(all_issues)}/{total} issues...")
            if len(all_issues) >= total or not issues:
                break
            start_at += len(issues)
        except Exception as e:
            print(f"[ERROR] {e}")
            break

    if not all_issues:
        print("[WARN] Không tìm thấy task nào thỏa mãn điều kiện!")
        return

    create_excel_report(all_issues, proj, output_file, client.base_url)

if __name__ == "__main__":
    main()
